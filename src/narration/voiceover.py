#!/usr/bin/env python3
"""
voiceover.py - narrate a slide deck from its speaker notes
==========================================================

By default it writes ONE self-contained web page: the slides, an audio player
per slide (play/pause, scrub bar, back/forward 10 s, speed), the transcript
(the speaker notes as written) beside the slide behind a draggable divider,
optional continuous play, and a Save button.  The MP3s are embedded, so the
page plays online and also works when downloaded and opened offline.
Optionally it also writes a PowerPoint with click-to-play audio on each slide.

Source (the deck to narrate), either:
  - a slides page built for the WPMW decks (HTML with <div id="stage">, one
    <section class="slide"> per slide and a `const NOTES = [...]` array), or
  - a .pptx: notes from its speaker notes, slide pictures rendered with
    LibreOffice (needs `soffice` and `pdftoppm`).

Output (in --outdir, default ./<name>_narration/):
  <name>_narrated.html   the narrated page
  audio/slide_NN.mp3     one clip per slide (NN = position, 01 = first)
  <name>_spoken.txt      the exact text sent to the voice, per slide
  <name>_narrated.pptx   only with --pptx

Voices (--tts):
  edge    Microsoft Edge neural voices (`pip install edge-tts`), needs internet. Default.
  piper   Piper neural voices (`pip install piper-tts`), run on this machine, offline
          after the voice is downloaded once (about 60 MB).  Steadier, less natural.
  flite   CMU Flite + ffmpeg, offline, robotic.
  silent  silent clips of the right length, for testing layout and plumbing.

Examples:
  python3 voiceover.py deck.html --pronunciations terms.json
  python3 voiceover.py deck.html --pronunciations terms.json --preview      # text only, no audio
  python3 voiceover.py deck.html --reuse-audio --reencode 32               # rebuild page, smaller file
  python3 voiceover.py deck.html --pptx exported.pptx                       # page + PowerPoint
  python3 voiceover.py talk.pptx --pptx                                     # pptx source, both outputs

Re-running is cheap: a slide whose spoken text, voice and rate are unchanged
keeps its clip (audio/slide_NN.txt records what it was made from), so after a
failure, or after fixing a pronunciation, only the affected slides are
synthesised again.  The Edge voice is asked for at most about 1200 characters
at a time; a request that stalls is abandoned after --timeout seconds, and a
failed one is retried three times, 5, 20 and 60 s apart.  --pause sets the
wait between requests (default 1 s); if many first tries are refused, a
longer pause (say 5 s) may avoid them at the cost of a slower run.  Exit status 3 means
some slides have no audio.

Clicks: each clip's joins between requests are listed in audio/slide_NN_joins.txt
(and in the log), so a click can be checked against them.  --max-chars 0 sends a
slide as one request (edge-tts itself still splits above about 4096 bytes);
--join fade joins the pieces with short fades; --declick runs ffmpeg's adeclick
on the page's copy of each clip, and --highpass 80 removes what lies below 80 Hz
from it (low thumps before words).  --resynth 14 (or 5,14) remakes only those slides.

Page keys: Left/Right change slide, Space or K play/pause, J/L back/forward
10 s, T show/hide transcript, C continuous play.  Drag the divider between
slide and transcript to resize it (double-click resets it).

Requirements: Python 3.9+.  edge-tts for --tts edge; ffmpeg for flite, silent
and --reencode; python-pptx and lxml for any .pptx input or output;
LibreOffice and poppler (pdftoppm) for a .pptx source.
"""

import argparse
import asyncio
import base64
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

# ==========================================================================
# 1. Text clean-up for speech
# ==========================================================================
BUILTIN_REPLACEMENTS = [
    ("√", " square root of "), ("∞", " infinity "), ("∇²", " del squared "),
    ("∇", " del "), ("∂", " partial "), ("→", " goes to "), ("⇒", " implies "),
    ("≈", " approximately "), ("≠", " is not equal to "),
    ("≤", " is less than or equal to "), ("≥", " is greater than or equal to "),
    ("±", " plus or minus "), ("×", " times "), ("·", " dot "), ("²", " squared"),
    ("³", " cubed"), ("—", ", "), ("–", ", "), ("‘", "'"), ("’", "'"), ("“", '"'),
    ("”", '"'), ("…", "..."), ("&", " and "), ("~", " approximately "),
]
SUBSCRIPTS = {"₀": "0", "₁": "1", "₂": "2", "₃": "3", "₄": "4", "₅": "5", "₆": "6",
              "₇": "7", "₈": "8", "₉": "9", "ₓ": "x"}


def greek_to_words(text):
    out = []
    for ch in text:
        if "Ͱ" <= ch <= "Ͽ" and ch.isalpha():
            name = unicodedata.name(ch, "")
            m = re.match(r"GREEK (?:SMALL|CAPITAL) LETTER ([A-Z]+)", name)
            if m:
                word = m.group(1).lower()
                out.append(f" {'capital ' if 'CAPITAL' in name else ''}{word} ")
                continue
        out.append(ch)
    return "".join(out)


def load_pronunciations(path):
    """JSON object {"written": "spoken"}, applied longest key first.

    Keys are literal text, matched anywhere (also inside words), so give a
    word that is pronounced by context (wind, live, read) enough surrounding
    words to be unambiguous: {"clock winds at": "clock wynds at"}.
    """
    if not path:
        return []
    with open(path, encoding="utf-8") as fh:
        data = json.load(fh)
    return sorted(data.items(), key=lambda kv: len(kv[0]), reverse=True)


def clean_for_tts(text, pronunciations=(), ascii_only=False):
    for written, spoken in pronunciations:
        text = text.replace(written, spoken)
    for old, new in BUILTIN_REPLACEMENTS:
        text = text.replace(old, new)
    text = greek_to_words(text)
    for old, new in SUBSCRIPTS.items():
        text = text.replace(old, " " + new)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    if ascii_only:
        text = re.sub(r"[^\x20-\x7E\n]", " ", text)
    return re.sub(r"[ \t]+", " ", text).strip()


# ==========================================================================
# 2. Sources
# ==========================================================================
class Deck:
    """name, notes[], head (HTML up to </head>), stage (the #stage div), size."""

    def __init__(self, name, notes, head, stage, size):
        self.name, self.notes, self.head, self.stage, self.size = name, notes, head, stage, size


def deck_from_html(path):
    src = Path(path).read_text(encoding="utf-8")
    m = re.search(r"const NOTES = (\[.*?\]);\n", src, re.S)
    if not m or '<div id="stage">' not in src or 'class="slide"' not in src:
        sys.exit(f"error: {path} is not a slides page of the expected kind "
                 "(needs <div id=\"stage\">, class=\"slide\" sections and a NOTES array)")
    if 'id="nv-audio"' in src:
        sys.exit(f"error: {path} is already a narrated page; give the plain slides page")
    notes = json.loads(m.group(1))
    start = src.index('<div id="stage">')
    end = src.index('<div id="notes"', start) if '<div id="notes"' in src[start:] else None
    if end is None:
        sys.exit("error: could not find the end of the #stage block (<div id=\"notes\"> missing)")
    stage = src[start:end].rstrip()
    n = stage.count('class="slide"')
    if n != len(notes):
        sys.exit(f"error: {n} slides but {len(notes)} notes in {path}")
    head = src[:src.index("</head>")]
    m = re.search(r"#stage\s*\{[^}]*?width:\s*(\d+)px;\s*height:\s*(\d+)px", src)
    size = (int(m.group(1)), int(m.group(2))) if m else (1920, 1080)
    return Deck(Path(path).stem, notes, head, stage, size)


def read_pptx_notes(pptx_path):
    from pptx import Presentation
    prs = Presentation(str(pptx_path))
    notes = []
    for slide in prs.slides:
        text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            text = slide.notes_slide.notes_text_frame.text.strip()
        notes.append(text)
    return notes, (int(prs.slide_width), int(prs.slide_height))


def deck_from_pptx(path, workdir):
    for tool in ("soffice", "pdftoppm"):
        if shutil.which(tool) is None:
            sys.exit(f"error: {tool} not found; a .pptx source needs LibreOffice and poppler-utils")
    notes, (w_emu, h_emu) = read_pptx_notes(path)
    width = 1920
    height = round(width * h_emu / w_emu)
    pdf_dir = Path(workdir) / "pdf"
    pdf_dir.mkdir()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(pdf_dir),
                    str(path)], check=True, capture_output=True, timeout=600)
    pdf = next(pdf_dir.glob("*.pdf"))
    subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(width), "-scale-to-y", str(height),
                    str(pdf), str(pdf_dir / "s")], check=True, capture_output=True, timeout=600)
    pngs = sorted(pdf_dir.glob("s-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
    if len(pngs) != len(notes):
        sys.exit(f"error: LibreOffice rendered {len(pngs)} pages for {len(notes)} slides")
    sections = []
    for k, png in enumerate(pngs, 1):
        b64 = base64.b64encode(png.read_bytes()).decode("ascii")
        sections.append(f'<section class="slide" data-n="{k}"><img alt="Slide {k}" '
                        f'src="data:image/png;base64,{b64}"></section>')
    title = html.escape(Path(path).stem.replace("_", " "))
    head = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
html, body {{ margin: 0; height: 100%; background: #1d232c; overflow: hidden; }}
#stage {{ position: absolute; left: 0; top: 0; width: {width}px; height: {height}px; transform-origin: 0 0; }}
.slide {{ position: absolute; left: 0; top: 0; width: {width}px; height: {height}px; display: block; }}
.slide img {{ width: 100%; height: 100%; display: block; }}
@media screen {{ .slide:not(.current) {{ display: none !important; }} }}
</style>
"""
    stage = '<div id="stage">\n' + "\n".join(sections) + "\n</div>"
    return Deck(Path(path).stem, notes, head, stage, (width, height))


# ==========================================================================
# 3. Speech synthesis
# ==========================================================================
EDGE_PIECE_CHARS = 1200        # characters per request to the voice service
EDGE_RETRY_WAITS = (5, 20, 60)  # seconds to wait before each retry of a piece
EDGE_TIMEOUT = 60              # seconds a single request may take (--timeout)
EDGE_PAUSE = 1.0               # seconds between requests (--pause)
EDGE_KBPS = 48                 # Edge sends 24 kHz mono MP3 at 48 kbit/s CBR
EDGE_RATE_HZ = 24000
JOIN_FADE_MS = 10              # fade at each join of two pieces (--join fade)


def split_for_speech(text, limit=EDGE_PIECE_CHARS):
    """Split at paragraph, then sentence, then word boundaries into pieces <= limit."""
    pieces = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if len(para) <= limit:
            pieces.append(para)
            continue
        cur = ""
        for sent in re.split(r"(?<=[.!?;:])\s+", para):
            while len(sent) > limit:                      # a very long sentence: cut at a space
                cut = sent.rfind(" ", 0, limit)
                cut = cut if cut > 0 else limit
                if cur:
                    pieces.append(cur)
                    cur = ""
                pieces.append(sent[:cut])
                sent = sent[cut:].lstrip()
            if cur and len(cur) + 1 + len(sent) > limit:
                pieces.append(cur)
                cur = sent
            else:
                cur = f"{cur} {sent}".strip()
        if cur:
            pieces.append(cur)
    # join short neighbours so the voice does not pause oddly between tiny requests
    merged = []
    for p in pieces:
        if merged and len(merged[-1]) + 2 + len(p) <= limit:
            merged[-1] += "\n\n" + p
        else:
            merged.append(p)
    return merged


async def _edge_bytes(text, voice, rate, timeout):
    import edge_tts

    async def collect():
        buf = bytearray()
        async for msg in edge_tts.Communicate(text, voice=voice, rate=rate).stream():
            if msg.get("type") == "audio":
                buf.extend(msg["data"])
        return bytes(buf)
    return await asyncio.wait_for(collect(), timeout)


def _decode_pcm(mp3_bytes, workdir, n):
    """Decode one piece to 16-bit mono PCM at the Edge rate (ffmpeg)."""
    src = Path(workdir) / f"piece_{n:02d}.mp3"
    src.write_bytes(mp3_bytes)
    r = subprocess.run(["ffmpeg", "-loglevel", "error", "-i", str(src), "-f", "s16le", "-ac", "1",
                        "-ar", str(EDGE_RATE_HZ), "-"], capture_output=True)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.decode(errors="replace").strip()[:200])
    return r.stdout


def join_pieces(datas, out_mp3, how="bytes"):
    """Write the pieces as one MP3.  Returns (join times in seconds, how they were joined).

    "bytes" appends the MP3 streams as they came (the old behaviour).  Each
    stream starts a fresh encoder, so the decoder can step at a join.  "fade"
    decodes each piece, fades JOIN_FADE_MS in and out at every join so the wave
    meets the join at zero, and encodes the whole clip once.  It needs ffmpeg;
    without it, or if decoding fails, the pieces are appended as bytes.
    """
    joins, t = [], 0.0
    for d in datas[:-1]:
        t += len(d) * 8 / (EDGE_KBPS * 1000)          # CBR: duration from size
        joins.append(t)
    tmp = out_mp3 + ".part"
    if how == "fade" and len(datas) > 1 and shutil.which("ffmpeg"):
        import array
        try:
            with tempfile.TemporaryDirectory() as work:
                pcm = [array.array("h", _decode_pcm(d, work, n)) for n, d in enumerate(datas, 1)]
                fade = EDGE_RATE_HZ * JOIN_FADE_MS // 1000
                joins, t = [], 0.0
                for n, a in enumerate(pcm):
                    m = min(fade, len(a) // 2)
                    if n > 0:                                  # fade in after a join
                        for i in range(m):
                            a[i] = int(a[i] * i / m)
                    if n < len(pcm) - 1:                       # fade out before a join
                        for i in range(m):
                            a[len(a) - 1 - i] = int(a[len(a) - 1 - i] * i / m)
                        t += len(a) / EDGE_RATE_HZ
                        joins.append(t)
                raw = b"".join(a.tobytes() for a in pcm)
                r = subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-f", "s16le", "-ar",
                                    str(EDGE_RATE_HZ), "-ac", "1", "-i", "-", "-b:a", "64k",
                                    "-f", "mp3", tmp], input=raw, capture_output=True)
                if r.returncode != 0:
                    raise RuntimeError(r.stderr.decode(errors="replace").strip()[:200])
            os.replace(tmp, out_mp3)
            return joins, "fade"
        except Exception as exc:
            print(f"\n      could not join with fades ({exc}); appending the pieces as bytes", end="")
    with open(tmp, "wb") as fh:
        for d in datas:
            fh.write(d)
    os.replace(tmp, out_mp3)
    return joins, "bytes"


def synth_edge(text, out_mp3, voice, rate, timeout=EDGE_TIMEOUT, pause=EDGE_PAUSE,
               max_chars=EDGE_PIECE_CHARS, join="bytes", **_):
    """Synthesise in pieces, each with a timeout and retries; write the MP3 only if all succeed.

    The Edge service sometimes stalls (no reply, no error) or answers "No audio
    was received" for a request it accepts at another moment.  Small requests,
    a timeout and a few spaced retries turn both into a short delay.
    max_chars = 0 sends the slide as one request (edge-tts itself still splits
    text longer than about 4096 bytes).  The join times are written beside the
    clip as slide_NN_joins.txt and printed, so a click can be checked against them.
    """
    pieces = split_for_speech(text, max_chars) if max_chars else [text]
    datas = []
    for n, piece in enumerate(pieces, 1):
        tag = f"piece {n}/{len(pieces)}"
        for attempt in range(len(EDGE_RETRY_WAITS) + 1):
            try:
                data = asyncio.run(_edge_bytes(piece, voice, rate, timeout))
                if not data:
                    raise RuntimeError("no audio received")
                datas.append(data)
                break
            except KeyboardInterrupt:
                raise
            except Exception as exc:  # timeout, network, service refusal ...
                why = "timed out" if isinstance(exc, asyncio.TimeoutError) else str(exc).split("\n")[0][:120]
                if attempt == len(EDGE_RETRY_WAITS):
                    print(f"\n      {tag}: {why}; giving up after {attempt + 1} tries", end="")
                    return False
                wait = EDGE_RETRY_WAITS[attempt]
                print(f"\n      {tag}: {why}; retry {attempt + 1} of {len(EDGE_RETRY_WAITS)} in {wait} s"
                      f" (then up to {timeout} s for a reply)", end="", flush=True)
                time.sleep(wait)
        time.sleep(pause)   # be gentle with the service between requests (--pause)
    joins, how = join_pieces(datas, out_mp3, join)
    note = Path(out_mp3).with_name(Path(out_mp3).stem + "_joins.txt")
    if joins:
        stamps = ", ".join(f"{int(j // 60)}:{j % 60:04.1f}" for j in joins)
        note.write_text(f"{len(pieces)} pieces, joined by {how}"
                        f" at {stamps}\n", encoding="utf-8")
        print(f"[{len(pieces)} pieces, {how} joins at {stamps}]", end="")
    elif note.exists():
        note.unlink()
    return os.path.getsize(out_mp3) > 0


def synth_flite(text, out_mp3, voice, **_):
    wav = out_mp3[:-4] + ".wav"
    r = subprocess.run(["flite", "-voice", voice, "-o", wav, "-t", text],
                       capture_output=True, timeout=600)
    if r.returncode != 0:
        print("\n      flite error:", r.stderr.decode(errors="replace")[:200])
        return False
    r = subprocess.run(["ffmpeg", "-y", "-i", wav, "-ar", "22050", "-ac", "1", "-b:a", "64k",
                        out_mp3], capture_output=True, timeout=600)
    if os.path.exists(wav):
        os.remove(wav)
    return r.returncode == 0 and os.path.exists(out_mp3)


def synth_silent(text, out_mp3, **_):
    seconds = max(1.0, len(text.split()) / 2.5)  # about 150 words per minute
    r = subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "anullsrc=r=22050:cl=mono",
                        "-t", f"{seconds:.2f}", "-b:a", "32k", out_mp3],
                       capture_output=True, timeout=120)
    return r.returncode == 0 and os.path.exists(out_mp3)


PIPER_VOICES_URL = "https://huggingface.co/rhasspy/piper-voices/resolve/main"


def piper_model(voice):
    """Path to a Piper voice: a .onnx file, or a name such as en_US-lessac-medium,
    downloaded once into ~/.cache/piper-voices/."""
    if voice.endswith(".onnx"):
        return Path(voice).expanduser()
    m = re.fullmatch(r"(([a-z]{2})_[A-Z]{2})-([A-Za-z0-9_]+)-(x_low|low|medium|high)", voice)
    if not m:
        sys.exit(f"error: piper voice {voice!r} is neither a .onnx file nor a name like en_US-lessac-medium")
    cache = Path.home() / ".cache" / "piper-voices"
    cache.mkdir(parents=True, exist_ok=True)
    model = cache / f"{voice}.onnx"
    for suffix in (".onnx", ".onnx.json"):
        target = cache / f"{voice}{suffix}"
        if not target.exists():
            url = f"{PIPER_VOICES_URL}/{m.group(2)}/{m.group(1)}/{m.group(3)}/{m.group(4)}/{voice}{suffix}"
            print(f"\n      downloading Piper voice {target.name} ...", end="", flush=True)
            import urllib.request
            tmp = target.with_suffix(target.suffix + ".part")
            urllib.request.urlretrieve(url, tmp)
            tmp.replace(target)
    return model


def synth_piper(text, out_mp3, voice, rate="+0%", **_):
    """Piper: a neural voice that runs on this machine, offline after the first download."""
    try:
        import piper  # noqa: F401
    except ImportError:
        sys.exit("error: Piper is not installed  ->  pip install piper-tts")
    m = re.fullmatch(r"([+-]\d+)%", rate or "+0%")
    length_scale = 1 / (1 + int(m.group(1)) / 100) if m else 1.0
    wav = out_mp3[:-4] + ".wav"
    r = subprocess.run([sys.executable, "-m", "piper", "-m", str(piper_model(voice)), "-f", wav,
                        "--length-scale", f"{length_scale:.3f}", "--sentence-silence", "0.25"],
                       input=text, text=True, capture_output=True, timeout=1800)
    if r.returncode != 0 or not os.path.exists(wav):
        print("\n      piper error:", (r.stderr or "")[-300:])
        return False
    r = subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", wav, "-ac", "1", "-b:a", "48k", out_mp3],
                       capture_output=True, timeout=600)
    os.remove(wav)
    return r.returncode == 0 and os.path.exists(out_mp3)


BACKENDS = {"edge": synth_edge, "piper": synth_piper, "flite": synth_flite, "silent": synth_silent}
DEFAULT_VOICE = {"edge": "en-US-AndrewNeural", "piper": "en_US-lessac-medium", "flite": "rms", "silent": ""}


def reencode_mp3(src, kbps, workdir, declick=False, highpass=0):
    """Re-encode a clip for the page, optionally filtered first.

    highpass > 0 removes what lies below that many Hz (ffmpeg highpass, two
    poles, 12 dB per octave): the low thump some voices put before a word.
    declick runs ffmpeg's adeclick, for sharp single-sample clicks.
    """
    out = Path(workdir) / ("re_" + Path(src).name)
    filters = ([f"highpass=f={highpass:g}"] if highpass else []) + (["adeclick"] if declick else [])
    r = subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-i", str(src), "-ac", "1"]
                       + (["-af", ",".join(filters)] if filters else [])
                       + ["-b:a", f"{kbps}k", str(out)], capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit(f"error: ffmpeg could not re-encode {src}: {r.stderr.strip()}")
    return out


# ==========================================================================
# 4. The narrated page
# ==========================================================================
VIEWER_CSS = r"""
<style id="narration-css">
html, body { overflow: hidden; }
#nv-side { position: fixed; top: 0; right: 0; bottom: 0; width: 34vw;
  background: #fbfaf6; color: #1d232c; display: flex; flex-direction: column;
  border-left: 3px solid #2F5D8A; font-family: 'IBM Plex Sans', Arial, sans-serif; box-sizing: border-box; }
#nv-side.hidden, #nv-split.hidden { display: none; }
#nv-side h2 { margin: 0; padding: 14px 20px 6px; font: 600 15px/1.3 'IBM Plex Sans', Arial, sans-serif;
  color: #2F5D8A; text-transform: uppercase; letter-spacing: 1px; }
#nv-text { flex: 1; overflow: auto; padding: 4px 20px 24px; font-size: 16px; line-height: 1.55;
  white-space: pre-wrap; }
#nv-split { position: fixed; z-index: 5; background: transparent; touch-action: none; }
#nv-split::after { content: ""; position: absolute; background: #8FB4DC; opacity: 0; transition: opacity .15s; }
#nv-split:hover::after, #nv-split:focus::after, body.nv-dragging #nv-split::after { opacity: .9; }
#nv-split.v { cursor: col-resize; } #nv-split.v::after { left: 3px; right: 3px; top: 0; bottom: 0; }
#nv-split.h { cursor: row-resize; } #nv-split.h::after { top: 3px; bottom: 3px; left: 0; right: 0; }
#nv-split:focus { outline: none; }
body.nv-dragging, body.nv-dragging * { user-select: none !important; }
#nv-bar { position: fixed; left: 0; bottom: 0; display: flex; flex-wrap: wrap; gap: 8px;
  align-items: center; justify-content: center; padding: 6px 10px; box-sizing: border-box;
  font: 14px 'IBM Plex Sans', Arial, sans-serif; color: #c9d1db; background: rgba(20,25,32,0.92); }
#nv-bar button, #nv-bar select, #nv-bar a { font: inherit; color: inherit; background: #2c3542;
  border: 1px solid #45505f; border-radius: 6px; padding: 4px 10px; cursor: pointer; text-decoration: none; }
#nv-bar button[aria-pressed="true"] { background: #2F5D8A; border-color: #8FB4DC; }
#nv-audio { height: 34px; width: min(420px, 38vw); }
#nv-count { min-width: 56px; text-align: center; }
@media print { #nv-side, #nv-bar, #nv-split { display: none !important; } }
</style>
"""

VIEWER_UI = r"""
<aside id="nv-side" aria-label="Transcript">
  <h2 id="nv-title">Transcript</h2>
  <div id="nv-text" tabindex="0"></div>
</aside>
<div id="nv-split" role="separator" tabindex="0" aria-label="Resize the transcript (drag, or arrow keys; double-click resets)"></div>
<div id="nv-bar">
  <button id="nv-prev" aria-label="Previous slide">&larr;</button>
  <span id="nv-count"></span>
  <button id="nv-next" aria-label="Next slide">&rarr;</button>
  <button id="nv-back" aria-label="Back 10 seconds">&minus;10 s</button>
  <audio id="nv-audio" controls preload="metadata"></audio>
  <button id="nv-fwd" aria-label="Forward 10 seconds">+10 s</button>
  <select id="nv-rate" aria-label="Playback speed">
    <option value="0.8">0.8&times;</option><option value="0.9">0.9&times;</option>
    <option value="1" selected>1&times;</option><option value="1.1">1.1&times;</option>
    <option value="1.25">1.25&times;</option><option value="1.5">1.5&times;</option>
  </select>
  <button id="nv-cont" aria-pressed="false" title="At the end of a clip, go on to the next slide (C)">Continuous</button>
  <button id="nv-tr" aria-pressed="true" title="Show or hide the transcript (T)">Transcript</button>
  <a id="nv-save" title="Save this page as one file that plays offline">Save</a>
</div>
"""

VIEWER_JS = r"""
<script>
const NOTES = __NOTES__;
const AUDIO = __AUDIO__;
const SW = __SW__, SH = __SH__, EMBEDDED = __EMBEDDED__, FILENAME = __FILENAME__;
const MISSING = __MISSING__;   // slides with notes whose audio could not be made
const slides = Array.from(document.querySelectorAll('.slide'));
const stage = document.getElementById('stage');
const side = document.getElementById('nv-side');
const split = document.getElementById('nv-split');
const bar = document.getElementById('nv-bar');
const audio = document.getElementById('nv-audio');
const text = document.getElementById('nv-text');
const contBtn = document.getElementById('nv-cont');
const trBtn = document.getElementById('nv-tr');
const rateSel = document.getElementById('nv-rate');
const save = document.getElementById('nv-save');
const store = { get(k, d) { try { const v = localStorage.getItem(k); return v === null ? d : v; } catch (e) { return d; } },
                set(k, v) { try { localStorage.setItem(k, v); } catch (e) {} } };
let i = -1;
let continuous = store.get('nv-cont', '0') === '1';
let showTr = store.get('nv-tr', '1') === '1';
let frac = parseFloat(store.get('nv-split', '0.34')) || 0.34;
rateSel.value = store.get('nv-rate', '1');
if (EMBEDDED) { save.href = location.href.split('#')[0]; save.setAttribute('download', FILENAME); }
else save.style.display = 'none';

const narrow = () => window.innerWidth <= 900;
const clampFrac = f => Math.max(0.15, Math.min(0.75, f));

function layout() {
  const nw = narrow();
  side.classList.toggle('hidden', !showTr);
  split.classList.toggle('hidden', !showTr);
  split.className = split.className.replace(/\b[vh]\b/g, '').trim();
  split.classList.add(nw ? 'h' : 'v');
  split.setAttribute('aria-orientation', nw ? 'horizontal' : 'vertical');
  split.setAttribute('aria-valuenow', String(Math.round(frac * 100)));
  if (nw) { side.style.width = ''; side.style.left = '0'; side.style.top = 'auto';
            side.style.height = (frac * 100) + 'vh'; side.style.borderLeft = 'none'; side.style.borderTop = '3px solid #2F5D8A'; }
  else    { side.style.height = ''; side.style.left = ''; side.style.top = '0';
            side.style.width = (frac * 100) + 'vw'; side.style.borderTop = 'none'; side.style.borderLeft = '3px solid #2F5D8A'; }
  const sideW = showTr && !nw ? side.getBoundingClientRect().width : 0;
  const sideH = showTr && nw ? side.getBoundingClientRect().height : 0;
  bar.style.right = sideW + 'px';
  bar.style.bottom = sideH + 'px';
  if (nw) Object.assign(split.style, { left: '0', right: '0', width: '', top: (window.innerHeight - sideH - 5) + 'px', bottom: '', height: '10px' });
  else Object.assign(split.style, { top: '0', bottom: '0', height: '', left: (window.innerWidth - sideW - 5) + 'px', right: '', width: '10px' });
  const barH = bar.getBoundingClientRect().height;
  const W = window.innerWidth - sideW, H = window.innerHeight - barH - sideH;
  const s = Math.max(0.05, Math.min(W / SW, H / SH));
  stage.style.transform = `translate(${(W - SW * s) / 2}px, ${(H - SH * s) / 2}px) scale(${s})`;
  trBtn.setAttribute('aria-pressed', String(showTr));
  contBtn.setAttribute('aria-pressed', String(continuous));
}
function show(k, play) {
  const n = Math.max(0, Math.min(slides.length - 1, k));
  const changed = n !== i;
  i = n;
  slides.forEach((el, j) => el.classList.toggle('current', j === i));
  document.getElementById('nv-count').textContent = `${i + 1} / ${slides.length}`;
  document.getElementById('nv-title').textContent = `Transcript, slide ${i + 1}`
    + (MISSING.includes(i + 1) ? ' (no audio yet)' : '');
  if (changed) {
    text.textContent = NOTES[i] || '(no notes)';
    text.scrollTop = 0;
    audio.pause();
    if (AUDIO[i]) { audio.src = AUDIO[i]; audio.style.visibility = 'visible'; }
    else { audio.removeAttribute('src'); audio.load(); audio.style.visibility = 'hidden'; }
    audio.playbackRate = parseFloat(rateSel.value);
  }
  if (play && AUDIO[i]) audio.play().catch(() => {});
  if (location.hash !== '#' + (i + 1)) history.replaceState(null, '', '#' + (i + 1));
}
function toggle() { if (!AUDIO[i]) return; audio.paused ? audio.play().catch(() => {}) : audio.pause(); }
function skip(d) { if (AUDIO[i]) audio.currentTime = Math.max(0, Math.min(audio.duration || 0, audio.currentTime + d)); }

// the draggable divider between slide and transcript
let dragging = false;
split.addEventListener('pointerdown', e => { dragging = true; split.setPointerCapture(e.pointerId); document.body.classList.add('nv-dragging'); e.preventDefault(); });
split.addEventListener('pointermove', e => {
  if (!dragging) return;
  frac = clampFrac(narrow() ? (window.innerHeight - e.clientY) / window.innerHeight
                            : (window.innerWidth - e.clientX) / window.innerWidth);
  layout();
});
const endDrag = () => { if (!dragging) return; dragging = false; document.body.classList.remove('nv-dragging'); store.set('nv-split', String(frac)); };
split.addEventListener('pointerup', endDrag);
split.addEventListener('pointercancel', endDrag);
split.addEventListener('dblclick', () => { frac = 0.34; store.set('nv-split', String(frac)); layout(); });
split.addEventListener('keydown', e => {
  const grow = narrow() ? 'ArrowUp' : 'ArrowLeft', shrink = narrow() ? 'ArrowDown' : 'ArrowRight';
  if (e.key === grow || e.key === shrink) {
    e.preventDefault(); e.stopPropagation();
    frac = clampFrac(frac + (e.key === grow ? 0.02 : -0.02)); store.set('nv-split', String(frac)); layout();
  }
});

audio.addEventListener('ended', () => { if (continuous && i < slides.length - 1) show(i + 1, true); });
audio.addEventListener('loadedmetadata', () => { audio.playbackRate = parseFloat(rateSel.value); });
rateSel.addEventListener('change', () => { audio.playbackRate = parseFloat(rateSel.value); store.set('nv-rate', rateSel.value); });
document.getElementById('nv-prev').onclick = () => show(i - 1, !audio.paused);
document.getElementById('nv-next').onclick = () => show(i + 1, !audio.paused);
document.getElementById('nv-back').onclick = () => skip(-10);
document.getElementById('nv-fwd').onclick = () => skip(10);
contBtn.onclick = () => { continuous = !continuous; store.set('nv-cont', continuous ? '1' : '0'); layout(); };
trBtn.onclick = () => { showTr = !showTr; store.set('nv-tr', showTr ? '1' : '0'); layout(); };
document.addEventListener('keydown', e => {
  if (e.target === split || (e.target.closest && e.target.closest('select'))) return;
  const k = e.key;
  if (['ArrowRight', 'PageDown'].includes(k)) { e.preventDefault(); show(i + 1, !audio.paused); }
  else if (['ArrowLeft', 'PageUp'].includes(k)) { e.preventDefault(); show(i - 1, !audio.paused); }
  else if (k === ' ' || k === 'k' || k === 'K') { e.preventDefault(); toggle(); }
  else if (k === 'j' || k === 'J') skip(-10);
  else if (k === 'l' || k === 'L') skip(10);
  else if (k === 't' || k === 'T') trBtn.click();
  else if (k === 'c' || k === 'C') contBtn.click();
  else if (k === 'Home') show(0, false);
  else if (k === 'End') show(slides.length - 1, false);
});
window.addEventListener('resize', layout);
window.addEventListener('hashchange', () => show((parseInt(location.hash.slice(1), 10) || 1) - 1, false));
layout();
show((parseInt(location.hash.slice(1), 10) || 1) - 1, false);
layout();
</script>
"""


def build_page(deck, audio_refs, out_path, embedded, missing=()):
    title = re.search(r"<title>(.*?)</title>", deck.head, re.S)
    head = deck.head
    if title:
        head = head.replace(title.group(0), f"<title>{title.group(1)} (narrated)</title>", 1)
    js = (VIEWER_JS
          .replace("__NOTES__", json.dumps(deck.notes, ensure_ascii=False).replace("</", "<\\/"))
          .replace("__AUDIO__", json.dumps(audio_refs))
          .replace("__SW__", str(deck.size[0])).replace("__SH__", str(deck.size[1]))
          .replace("__EMBEDDED__", "true" if embedded else "false")
          .replace("__MISSING__", json.dumps(sorted(missing)))
          .replace("__FILENAME__", json.dumps(Path(out_path).name)))
    page = head + VIEWER_CSS + "</head>\n<body>\n" + deck.stage + "\n" + VIEWER_UI + js + "</body>\n</html>\n"
    Path(out_path).write_text(page, encoding="utf-8")


# ==========================================================================
# 5. PowerPoint output: click-to-play audio on each slide
# ==========================================================================
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main"}
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
AUDIO_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio"
MEDIA_REL = "http://schemas.microsoft.com/office/2007/relationships/media"
EMU_PER_INCH = 914400


def inches(x):
    return int(x * EMU_PER_INCH)


def slide_files_in_order(pptx_path):
    """Slide XML files in presentation order (slideN.xml is NOT the Nth slide after reordering)."""
    from lxml import etree
    with zipfile.ZipFile(pptx_path) as z:
        pres = etree.fromstring(z.read("ppt/presentation.xml"))
        rels = etree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels}
    files = []
    for sld in pres.iterfind(f".//{{{NS['p']}}}sldId"):
        t = target[sld.get(f"{{{NS['r']}}}id")]
        files.append(t.lstrip("/") if t.startswith("/") else "ppt/" + t)
    return files


def _pic_xml(shape_id, name, audio_rid, media_rid, x, y, size):
    # action="ppaction://media" makes the shape a play/pause toggle in Slide Show
    return f"""<p:pic xmlns:p="{NS['p']}" xmlns:a="{NS['a']}" xmlns:r="{NS['r']}">
  <p:nvPicPr>
    <p:cNvPr id="{shape_id}" name="{name}"><a:hlinkClick r:id="" action="ppaction://media"/></p:cNvPr>
    <p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>
    <p:nvPr>
      <a:audioFile r:link="{audio_rid}"/>
      <p:extLst><p:ext uri="{{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}}">
        <p14:media xmlns:p14="{NS['p14']}" r:embed="{media_rid}"/></p:ext></p:extLst>
    </p:nvPr>
  </p:nvPicPr>
  <p:blipFill><a:blip/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
  <p:spPr>
    <a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{size}" cy="{size}"/></a:xfrm>
    <a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom>
    <a:solidFill><a:srgbClr val="06B6D4"><a:alpha val="85000"/></a:srgbClr></a:solidFill>
    <a:ln w="12700"><a:solidFill><a:srgbClr val="FFFFFF"><a:alpha val="60000"/></a:srgbClr></a:solidFill></a:ln>
  </p:spPr>
</p:pic>"""


def _label_xml(shape_id, name, text, x, y, w, h, color):
    return f"""<p:sp xmlns:p="{NS['p']}" xmlns:a="{NS['a']}" xmlns:r="{NS['r']}">
  <p:nvSpPr>
    <p:cNvPr id="{shape_id}" name="{name}"><a:hlinkClick r:id="" action="ppaction://media"/></p:cNvPr>
    <p:cNvSpPr txBox="1"/><p:nvPr/>
  </p:nvSpPr>
  <p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>
  <p:txBody>
    <a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/><a:lstStyle/>
    <a:p><a:pPr algn="r"/><a:r><a:rPr lang="en-US" sz="1000" dirty="0">
      <a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:rPr><a:t>{escape(text)}</a:t></a:r></a:p>
  </p:txBody>
</p:sp>"""


def embed_audio(pptx_path, audio_files, output_path, label="▶ Narration", label_color="94A3B8"):
    from lxml import etree
    _, (slide_w, slide_h) = read_pptx_notes(pptx_path)
    slide_parts = slide_files_in_order(pptx_path)
    icon, margin = inches(0.4), inches(0.25)
    ix, iy = slide_w - margin - icon, slide_h - margin - icon
    work = Path(tempfile.mkdtemp(prefix="voiceover_pptx_"))
    try:
        with zipfile.ZipFile(pptx_path) as zin:
            zin.extractall(work)
        (work / "ppt" / "media").mkdir(parents=True, exist_ok=True)
        ct_path = work / "[Content_Types].xml"
        ct_tree = etree.parse(str(ct_path))
        ct_root = ct_tree.getroot()
        if not any(e.get("Extension", "").lower() == "mp3" for e in ct_root.findall(f"{{{CT_NS}}}Default")):
            d = etree.SubElement(ct_root, f"{{{CT_NS}}}Default")
            d.set("Extension", "mp3")
            d.set("ContentType", "audio/mpeg")
            ct_root.insert(0, d)
        ct_tree.write(str(ct_path), xml_declaration=True, encoding="UTF-8", standalone=True)
        for idx, audio in enumerate(audio_files):
            if not audio or idx >= len(slide_parts):
                continue
            num = idx + 1
            slide_xml = work / slide_parts[idx]
            rels_file = slide_xml.parent / "_rels" / (slide_xml.name + ".rels")
            rels_file.parent.mkdir(exist_ok=True)
            media_name = f"voiceover_{num:02d}.mp3"
            shutil.copyfile(audio, work / "ppt" / "media" / media_name)
            if rels_file.exists():
                rels_tree = etree.parse(str(rels_file))
                rels_root = rels_tree.getroot()
            else:
                rels_root = etree.Element(f"{{{REL_NS}}}Relationships", nsmap={None: REL_NS})
                rels_tree = etree.ElementTree(rels_root)
            used = [int(m.group(1)) for el in rels_root if (m := re.fullmatch(r"rId(\d+)", el.get("Id", "")))]
            nxt = max(used, default=0) + 1
            audio_rid, media_rid = f"rId{nxt}", f"rId{nxt + 1}"
            for rid, rtype in ((audio_rid, AUDIO_REL), (media_rid, MEDIA_REL)):
                rel = etree.SubElement(rels_root, f"{{{REL_NS}}}Relationship")
                rel.set("Id", rid)
                rel.set("Type", rtype)
                rel.set("Target", f"../media/{media_name}")
            rels_tree.write(str(rels_file), xml_declaration=True, encoding="UTF-8", standalone=True)
            tree = etree.parse(str(slide_xml))
            sp_tree = tree.getroot().find(f".//{{{NS['p']}}}spTree")
            if sp_tree is None:
                print(f"  slide {num}: no shape tree, skipped")
                continue
            ids = [int(e.get("id")) for e in tree.getroot().iter(f"{{{NS['p']}}}cNvPr") if (e.get("id") or "").isdigit()]
            sid = max(ids, default=1) + 1
            sp_tree.append(etree.fromstring(_pic_xml(sid, f"Voice-over {num}", audio_rid, media_rid, ix, iy, icon)))
            if label:
                lw, lh = inches(1.3), inches(0.3)
                sp_tree.append(etree.fromstring(_label_xml(sid + 1, f"Voice-over label {num}", label,
                                                           ix - lw - inches(0.05), iy + (icon - lh) // 2,
                                                           lw, lh, label_color)))
            tree.write(str(slide_xml), xml_declaration=True, encoding="UTF-8", standalone=True)
        files = [p for p in work.rglob("*") if p.is_file()]
        files.sort(key=lambda p: (p.relative_to(work).as_posix() != "[Content_Types].xml", p.relative_to(work).as_posix()))
        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for p in files:
                zout.write(p, p.relative_to(work).as_posix())
    finally:
        shutil.rmtree(work, ignore_errors=True)


# ==========================================================================
# 6. Main
# ==========================================================================
def main():
    ap = argparse.ArgumentParser(description="Narrate a slide deck: a web page by default, PowerPoint optionally.",
                                 formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__.split("Examples:")[1])
    ap.add_argument("source", help="slides page (.html) or deck (.pptx) whose speaker notes are narrated")
    ap.add_argument("-d", "--outdir", help="output folder (default: ./<name>_narration)")
    ap.add_argument("--name", help="base name of the outputs (default: the source's file name)")
    ap.add_argument("--pptx", nargs="?", const="", metavar="DECK.pptx",
                    help="also write a narrated PowerPoint; give the exported deck for an .html source")
    ap.add_argument("--no-html", action="store_true", help="do not write the web page")
    ap.add_argument("--tts", choices=sorted(BACKENDS), default="edge")
    ap.add_argument("--voice", help="voice (edge default en-US-AndrewNeural; flite default rms)")
    ap.add_argument("--rate", default="+0%", help="edge speaking rate, e.g. '+8%%' or '-5%%'")
    ap.add_argument("--pronunciations", help="JSON {written: spoken} applied before synthesis")
    ap.add_argument("--preview", action="store_true", help="write <name>_spoken.txt and stop (no audio)")
    ap.add_argument("--reuse-audio", action="store_true",
                    help="no synthesis at all: build from the MP3s already in <outdir>/audio")
    ap.add_argument("--resynth", nargs="?", const="all", metavar="SLIDES",
                    help="synthesise again even if unchanged: every slide, or only the listed ones, "
                         "e.g. --resynth 14 or --resynth 5,14")
    ap.add_argument("--max-chars", type=int, default=EDGE_PIECE_CHARS, metavar="N",
                    help=f"edge: characters per request (default {EDGE_PIECE_CHARS}); "
                         "0 sends each slide as one request")
    ap.add_argument("--join", choices=("bytes", "fade"), default="bytes",
                    help="edge: how pieces are joined: appended as they came (default), or decoded, "
                         f"faded over {JOIN_FADE_MS} ms at each join and encoded once (needs ffmpeg)")
    ap.add_argument("--declick", action="store_true",
                    help="run ffmpeg's adeclick filter on each clip for the page (the clips in "
                         "audio/ are untouched, so this works with --reuse-audio)")
    ap.add_argument("--highpass", type=float, default=0, metavar="HZ",
                    help="remove what lies below HZ from the page's copy of each clip, e.g. 80, "
                         "to soften low thumps before words (works with --reuse-audio)")
    ap.add_argument("--pause", type=float, default=EDGE_PAUSE, metavar="SECONDS",
                    help=f"edge: seconds to wait between requests (default {EDGE_PAUSE:g}); "
                         "a longer pause, e.g. 5, can avoid the refusals that need a retry")
    ap.add_argument("--timeout", type=int, default=EDGE_TIMEOUT,
                    help=f"edge: seconds one request may take before it is retried (default {EDGE_TIMEOUT})")
    ap.add_argument("--reencode", type=int, metavar="KBPS", help="re-encode clips to KBPS kbit/s mono for the page (ffmpeg)")
    ap.add_argument("--link", action="store_true", help="link audio/slide_NN.mp3 instead of embedding it in the page")
    ap.add_argument("--label", default="▶ Narration", help="PowerPoint: text beside the play icon ('' for none)")
    ap.add_argument("--label-color", default="94A3B8", help="PowerPoint: label colour RRGGBB")
    args = ap.parse_args()
    if args.resynth in (None, "all"):
        resynth = set() if args.resynth is None else None      # None = every slide
    else:
        try:
            resynth = {int(x) for x in args.resynth.replace(" ", "").split(",") if x}
        except ValueError:
            sys.exit(f"error: --resynth takes slide numbers such as 5,14 (got {args.resynth!r})")

    src = Path(args.source).expanduser()
    if not src.is_file():
        sys.exit(f"error: {src} not found")
    is_pptx = src.suffix.lower() == ".pptx"
    pptx_in = None
    if args.pptx is not None:
        pptx_in = src if (args.pptx == "" and is_pptx) else (Path(args.pptx).expanduser() if args.pptx else None)
        if pptx_in is None:
            sys.exit("error: --pptx needs the exported .pptx when the source is a web page")
        if not pptx_in.is_file():
            sys.exit(f"error: {pptx_in} not found")

    work = tempfile.mkdtemp(prefix="voiceover_")
    try:
        need_page = not args.no_html
        if is_pptx:
            deck = deck_from_pptx(src, work) if need_page else Deck(src.stem, read_pptx_notes(src)[0], "", "", (0, 0))
        else:
            deck = deck_from_html(src)
        name = args.name or deck.name
        outdir = Path(args.outdir).expanduser() if args.outdir else Path.cwd() / f"{name}_narration"
        audio_dir = outdir / "audio"
        audio_dir.mkdir(parents=True, exist_ok=True)
        if pptx_in is not None:
            pnotes, _ = read_pptx_notes(pptx_in)
            if len(pnotes) != len(deck.notes):
                sys.exit(f"error: {pptx_in} has {len(pnotes)} slides, the source has {len(deck.notes)}")

        # text that will be spoken
        pron = load_pronunciations(args.pronunciations)
        spoken = [clean_for_tts(n, pron, ascii_only=(args.tts == "flite")) if n else "" for n in deck.notes]
        words = [len(s.split()) for s in spoken]
        with open(outdir / f"{name}_spoken.txt", "w", encoding="utf-8") as fh:
            for k, s in enumerate(spoken, 1):
                left = sorted(set(re.findall(r"[^\x00-\x7f]", s)))
                fh.write(f"--- slide {k:02d}: {words[k-1]} words, about {words[k-1]/150:.1f} min"
                         + (f"; NOT SPOKEN WELL: {' '.join(left)}" if left else "") + " ---\n" + s + "\n\n")
        print(f"source : {src}  ({len(deck.notes)} slides, {sum(bool(n) for n in deck.notes)} with notes)")
        print(f"outdir : {outdir}")
        print(f"text   : {outdir / (name + '_spoken.txt')}  ({sum(words)} words, about {sum(words)/150:.0f} min)")
        unspoken = sorted({c for s in spoken for c in re.findall(r"[^\x00-\x7f]", s)})
        if unspoken:
            print(f"warning: characters the voice may not read well: {' '.join(unspoken)} "
                  "(see NOT SPOKEN WELL in the text file; add them to the pronunciations)")
        if args.preview:
            return

        # audio
        voice = args.voice or DEFAULT_VOICE[args.tts]
        if not args.reuse_audio:
            if args.tts == "edge":
                try:
                    import edge_tts  # noqa: F401
                except ImportError:
                    sys.exit("error: edge-tts is not installed  ->  pip install edge-tts")
            if args.tts == "piper":
                try:
                    import piper  # noqa: F401
                except ImportError:
                    sys.exit("error: Piper is not installed  ->  pip install piper-tts")
            missing = [t for t in {"flite": ["flite", "ffmpeg"], "silent": ["ffmpeg"],
                                   "piper": ["ffmpeg"]}.get(args.tts, [])
                       if shutil.which(t) is None]
            if missing:
                sys.exit(f"error: missing tool(s): {', '.join(missing)}")
        if (args.reencode or args.declick or args.highpass) and shutil.which("ffmpeg") is None:
            sys.exit("error: --reencode, --declick and --highpass need ffmpeg")
        print(f"audio  : {audio_dir}  ({'reusing' if args.reuse_audio else args.tts + ' ' + voice})", flush=True)
        clips, failed, kept = [], [], 0
        run_of_failures, gave_up = 0, False
        for k, text in enumerate(spoken, 1):
            mp3 = audio_dir / f"slide_{k:02d}.mp3"
            stamp = audio_dir / f"slide_{k:02d}.txt"        # what the clip was made from
            key = f"{args.tts}|{voice}|{args.rate}\n{text}"
            if args.reuse_audio:
                ok = mp3.exists() and mp3.stat().st_size > 0
                print(f"    slide {k:2d}: {'found' if ok else 'no mp3'}", flush=True)
            elif not text:
                print(f"    slide {k:2d}: no notes", flush=True)
                ok = False
                for f in (mp3, stamp):
                    if f.exists():
                        f.unlink()
            elif (resynth is not None and k not in resynth and mp3.exists() and mp3.stat().st_size > 0 and stamp.exists()
                  and stamp.read_text(encoding="utf-8") == key):
                print(f"    slide {k:2d}: unchanged, kept ({mp3.stat().st_size // 1024} KB)", flush=True)
                ok = True
                kept += 1
            elif gave_up:
                print(f"    slide {k:2d}: not attempted", flush=True)
                ok = False
            else:
                print(f"    slide {k:2d}: synthesising ... ", end="", flush=True)
                ok = BACKENDS[args.tts](text, str(mp3), voice=voice, rate=args.rate,
                                        timeout=args.timeout, pause=args.pause,
                                        max_chars=args.max_chars, join=args.join)
                if ok:
                    stamp.write_text(key, encoding="utf-8")
                elif stamp.exists():
                    stamp.unlink()                       # an old clip no longer matches the text
                print(f" ok ({mp3.stat().st_size // 1024} KB)" if ok else " FAILED", flush=True)
            if not ok and text:
                failed.append(k)
            if text and not args.reuse_audio and not gave_up:
                run_of_failures = 0 if ok else run_of_failures + 1
                if run_of_failures >= 2:
                    gave_up = True
                    print(f"\n    Two slides in a row failed after every retry: the {args.tts} service is"
                          " refusing this machine for now (Edge does this after heavy use). Stopping;"
                          " clips made so far are kept.", flush=True)
            clips.append(mp3 if ok else None)
        if not any(clips):
            if failed and args.tts == "edge":
                print("\nNo audio could be made. Check the service with  edge-tts --text 'test' --write-media"
                      " /tmp/t.mp3 , update with  pip install -U edge-tts , wait an hour or two and run again,"
                      " or use the offline voice:  --tts piper")
            sys.exit(3 if failed else "nothing to narrate (no notes)")
        if kept:
            print(f"         {kept} clip(s) kept from an earlier run (--resynth makes them again)")

        # page
        if need_page:
            refs = []
            for c in clips:
                if c is None:
                    refs.append(None)
                    continue
                redo = bool(args.reencode or args.declick or args.highpass)
                f = (reencode_mp3(c, args.reencode or EDGE_KBPS, work, args.declick, args.highpass)
                     if redo else c)
                if args.link:
                    sub = "audio_web" if redo else "audio"
                    if redo:  # keep the originals in audio/ untouched
                        (outdir / sub).mkdir(exist_ok=True)
                        shutil.copyfile(f, outdir / sub / c.name)
                    refs.append(f"{sub}/{c.name}")
                else:
                    refs.append("data:audio/mpeg;base64," + base64.b64encode(Path(f).read_bytes()).decode("ascii"))
            page = outdir / f"{name}_narrated.html"
            build_page(deck, refs, page, embedded=not args.link, missing=failed)
            print(f"page   : {page}  ({page.stat().st_size / 1e6:.1f} MB, audio "
                  f"{'linked' if args.link else 'embedded'})")

        # PowerPoint
        if pptx_in is not None:
            out_pptx = outdir / f"{Path(pptx_in).stem}_narrated.pptx"
            embed_audio(str(pptx_in), [str(c) if c else None for c in clips], str(out_pptx),
                        label=args.label, label_color=args.label_color)
            print(f"pptx   : {out_pptx}  (click the play icon in Slide Show)")
        if failed:
            print(f"\nMISSING AUDIO for slide(s) {', '.join(map(str, failed))}. The page was written, but do not"
                  f" publish it yet: run the same command again later; slides already made are kept, and only"
                  f" the missing ones are synthesised.")
            if args.tts == "edge":
                print("If Edge keeps refusing: test it with  edge-tts --text 'test' --write-media /tmp/t.mp3 ,"
                      " update it with  pip install -U edge-tts , wait an hour or two, or use the offline"
                      " voice for the missing slides with  --tts piper  (a different voice from the rest).")
            sys.exit(3)
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
